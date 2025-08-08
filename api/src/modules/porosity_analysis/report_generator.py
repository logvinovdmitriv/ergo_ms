import os
import logging
from datetime import datetime
from typing import Dict, List, Optional
from PIL import Image
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image as RLImage, PageBreak
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.fonts import addMapping
import matplotlib
matplotlib.use('Agg')

# Регистрируем шрифт с поддержкой кириллицы
def setup_russian_fonts():
    """Настраивает шрифты с поддержкой кириллицы для PDF"""
    try:
        # Проверяем переменную окружения для пользовательского шрифта
        custom_font_path = os.environ.get('POROSITY_PDF_FONT_PATH')
        if custom_font_path and os.path.exists(custom_font_path):
            try:
                font_name = os.path.splitext(os.path.basename(custom_font_path))[0]
                pdfmetrics.registerFont(TTFont(font_name, custom_font_path))
                addMapping(font_name, 0, 0, font_name)
                print(f"Используется пользовательский шрифт: {font_name} из {custom_font_path}")
                return font_name
            except Exception as font_error:
                print(f"Не удалось зарегистрировать пользовательский шрифт {custom_font_path}: {font_error}")
        
        # Список возможных путей к шрифтам с поддержкой кириллицы
        font_paths = [
            # Linux пути
            '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',
            '/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf',
            '/usr/share/fonts/truetype/freefont/FreeSans.ttf',
            '/usr/share/fonts/TTF/DejaVuSans.ttf',
            '/usr/share/fonts/TTF/LiberationSans-Regular.ttf',
            
            # Windows пути (если запускается на Windows)
            'C:/Windows/Fonts/arial.ttf',
            'C:/Windows/Fonts/calibri.ttf',
            'C:/Windows/Fonts/tahoma.ttf',
            
            # macOS пути
            '/System/Library/Fonts/Arial.ttf',
            '/System/Library/Fonts/Helvetica.ttc',
            
            # Альтернативные Linux пути
            '/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf',
            '/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf',
            '/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf',
        ]
        
        for font_path in font_paths:
            if os.path.exists(font_path):
                try:
                    font_name = os.path.splitext(os.path.basename(font_path))[0]
                    pdfmetrics.registerFont(TTFont(font_name, font_path))
                    addMapping(font_name, 0, 0, font_name)
                    print(f"Успешно зарегистрирован шрифт: {font_name} из {font_path}")
                    return font_name
                except Exception as font_error:
                    print(f"Не удалось зарегистрировать шрифт {font_path}: {font_error}")
                    continue
        
        # Если ничего не найдено, используем встроенный шрифт
        print("Не найдены шрифты с поддержкой кириллицы, используем стандартный Helvetica")
        return 'Helvetica'
        
    except Exception as e:
        print(f"Ошибка при настройке шрифтов: {e}")
        return 'Helvetica'

# Инициализируем шрифт при импорте модуля
RUSSIAN_FONT = setup_russian_fonts()

# Выводим информацию о выбранном шрифте
print(f"Используется шрифт для PDF: {RUSSIAN_FONT}")

# Настраиваем логгер для генератора отчетов
logger = logging.getLogger('celery.task.porosity_analysis.reports')


class PorosityReportGenerator:
    """Генератор отчетов по анализу пористости"""
    
    def __init__(self, analysis):
        """
        Инициализация генератора отчетов
        
        Args:
            analysis: Объект PorosityAnalysis из базы данных
        """
        self.analysis = analysis
        self.results_dir = analysis.results_directory
        
    def _get_available_images(self) -> List[tuple]:
        """
        Получает список доступных изображений для отчета
        
        Returns:
            List[tuple]: Список кортежей (путь_к_файлу, описание)
        """
        available_images = []
        
        # Список изображений для включения в отчет с их описаниями
        images_to_check = [
            ('original_image_path', 'Исходное изображение', self.analysis.original_image_path),
            ('image_with_scale_bar.png', 'Изображение с обнаруженной шкалой', os.path.join(self.results_dir, 'image_with_scale_bar.png')),
            ('figure1_contrast.png', 'Этапы обработки контраста', os.path.join(self.results_dir, 'figure1_contrast.png')),
            ('figure2_excluded_areas.png', 'Исключенные области', os.path.join(self.results_dir, 'figure2_excluded_areas.png')),
            ('figure3_texture_clusters.png', 'Текстурный анализ', os.path.join(self.results_dir, 'figure3_texture_clusters.png')),
            ('figure4_mask_result.png', 'Бинарная маска и результат анализа', os.path.join(self.results_dir, 'figure4_mask_result.png')),
            ('figure5_overlay.png', 'Наложение результатов на исходное изображение', os.path.join(self.results_dir, 'figure5_overlay.png')),
            ('pore_size_distribution.png', 'Распределение размеров пор', os.path.join(self.results_dir, 'pore_size_distribution.png')),
            ('interpore_distances.png', 'Межпоровые расстояния', os.path.join(self.results_dir, 'interpore_distances.png')),
            ('pore_orientation_rose.png', 'Роза направлений пор', os.path.join(self.results_dir, 'pore_orientation_rose.png')),
            ('pore_orientation_histogram.png', 'Гистограмма ориентации пор', os.path.join(self.results_dir, 'pore_orientation_histogram.png')),
            ('pore_shapes_analysis.png', 'Анализ форм пор', os.path.join(self.results_dir, 'pore_shapes_analysis.png')),
            ('circularity_distribution.png', 'Распределение кругового фактора', os.path.join(self.results_dir, 'circularity_distribution.png')),
            ('ellipticity_vs_area.png', 'Зависимость эллиптичности от площади', os.path.join(self.results_dir, 'ellipticity_vs_area.png'))
        ]
        
        for image_name, description, image_path in images_to_check:
            if os.path.exists(image_path) and os.path.isfile(image_path):
                # Проверяем, что файл не пустой
                if os.path.getsize(image_path) > 0:
                    available_images.append((image_path, description))
                    logger.info(f"Найдено изображение для отчета: {image_name} -> {image_path}")
                else:
                    logger.warning(f"Файл изображения пустой: {image_path}")
            else:
                logger.warning(f"Файл изображения не найден: {image_path}")
        
        return available_images
        
    def generate_docx_report(self, output_path: str) -> bool:
        """
        Генерирует отчет в формате .docx
        
        Args:
            output_path: Путь для сохранения файла отчета
            
        Returns:
            bool: True если отчет создан успешно
        """
        try:
            # Создаем документ
            doc = Document()
            
            # Заголовок
            title = doc.add_heading('Отчет по анализу пористости', 0)
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            # Информация об анализе
            doc.add_heading('Общая информация', level=1)
            info_table = doc.add_table(rows=4, cols=2)
            info_table.style = 'Table Grid'
            
            info_data = [
                ('Название анализа:', self.analysis.name),
                ('Дата проведения:', self.analysis.created_at.strftime('%d.%m.%Y %H:%M')),
                ('Статус:', 'Завершен' if self.analysis.status == 'completed' else self.analysis.status),
                ('Масштаб:', f'{self.analysis.scale_value} мкм')
            ]
            
            for i, (label, value) in enumerate(info_data):
                info_table.cell(i, 0).text = label
                info_table.cell(i, 1).text = str(value)
                # Делаем первый столбец жирным
                info_table.cell(i, 0).paragraphs[0].runs[0].bold = True
            
            doc.add_paragraph()
            
            # Результаты анализа
            doc.add_heading('Результаты анализа', level=1)
            
            # Основные показатели
            doc.add_heading('Основные показатели', level=2)
            main_table = doc.add_table(rows=7, cols=2)
            main_table.style = 'Table Grid'
            
            main_data = [
                ('Процент пористости:', f'{self.analysis.porosity_percentage:.2f}%' if self.analysis.porosity_percentage else 'Н/Д'),
                ('Количество пор:', str(self.analysis.number_of_pores) if self.analysis.number_of_pores else 'Н/Д'),
                ('Средний размер пор:', f'{self.analysis.average_pore_size:.2f} мкм' if self.analysis.average_pore_size else 'Н/Д'),
                ('Максимальный размер пор:', f'{self.analysis.max_pore_size:.2f} мкм' if self.analysis.max_pore_size else 'Н/Д'),
                ('Минимальный размер пор:', f'{self.analysis.min_pore_size:.2f} мкм' if self.analysis.min_pore_size else 'Н/Д'),
                ('Плотность пор:', f'{self.analysis.pore_density:.4f} пор/мкм²' if self.analysis.pore_density else 'Н/Д'),
                ('Среднее межпоровое расстояние:', f'{self.analysis.average_interpore_distance:.2f} мкм' if self.analysis.average_interpore_distance else 'Н/Д')
            ]
            
            for i, (label, value) in enumerate(main_data):
                main_table.cell(i, 0).text = label
                main_table.cell(i, 1).text = value
                main_table.cell(i, 0).paragraphs[0].runs[0].bold = True
            
            # Добавляем изображения
            doc.add_page_break()
            doc.add_heading('Результаты визуализации', level=1)
            
            # Получаем список доступных изображений
            available_images = self._get_available_images()
            
            if not available_images:
                doc.add_paragraph('Изображения результатов анализа не найдены.')
            else:
                for image_path, description in available_images:
                    doc.add_heading(description, level=2)
                    try:
                        # Добавляем изображение с ограничением размера
                        doc.add_picture(image_path, width=Inches(5.5))
                        doc.add_paragraph()
                    except Exception as e:
                        logger.warning(f"Не удалось добавить изображение {image_path}: {e}")
                        doc.add_paragraph(f"[Изображение {description} недоступно]")
            
            # Сохраняем документ
            doc.save(output_path)
            logger.info(f"Отчет DOCX сохранен: {output_path}")
            return True
            
        except Exception as e:
            logger.error(f"Ошибка при генерации DOCX отчета: {e}")
            return False
    
    def generate_pdf_report(self, output_path: str) -> bool:
        """
        Генерирует отчет в формате .pdf
        
        Args:
            output_path: Путь для сохранения файла отчета
            
        Returns:
            bool: True если отчет создан успешно
        """
        try:
            # Создаем документ PDF
            doc = SimpleDocTemplate(output_path, pagesize=A4)
            story = []
            
            # Стили с поддержкой кириллицы
            styles = getSampleStyleSheet()
            title_style = ParagraphStyle(
                'CustomTitle',
                parent=styles['Heading1'],
                fontName=RUSSIAN_FONT,
                fontSize=24,
                textColor=colors.HexColor('#1a1a1a'),
                alignment=TA_CENTER,
                spaceAfter=30
            )
            
            heading_style = ParagraphStyle(
                'CustomHeading',
                parent=styles['Heading2'],
                fontName=RUSSIAN_FONT,
                fontSize=16,
                textColor=colors.HexColor('#333333'),
                spaceAfter=12
            )
            
            # Создаем стиль для обычного текста с поддержкой кириллицы
            normal_style = ParagraphStyle(
                'CustomNormal',
                parent=styles['Normal'],
                fontName=RUSSIAN_FONT,
                fontSize=12,
                textColor=colors.HexColor('#000000')
            )
            
            # Заголовок
            story.append(Paragraph('Отчет по анализу пористости', title_style))
            story.append(Spacer(1, 0.2*inch))
            
            # Информация об анализе
            story.append(Paragraph('Общая информация', heading_style))
            
            info_data = [
                ['Параметр', 'Значение'],
                ['Название анализа:', self.analysis.name],
                ['Дата проведения:', self.analysis.created_at.strftime('%d.%m.%Y %H:%M')],
                ['Статус:', 'Завершен' if self.analysis.status == 'completed' else self.analysis.status],
                ['Масштаб:', f'{self.analysis.scale_value} мкм']
            ]
            
            info_table = Table(info_data, colWidths=[3*inch, 3*inch])
            info_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('FONTNAME', (0, 0), (-1, 0), RUSSIAN_FONT),
                ('FONTSIZE', (0, 0), (-1, 0), 12),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
                ('FONTNAME', (0, 1), (-1, -1), RUSSIAN_FONT),
                ('FONTSIZE', (0, 1), (-1, -1), 11),
                ('GRID', (0, 0), (-1, -1), 1, colors.black)
            ]))
            
            story.append(info_table)
            story.append(Spacer(1, 0.3*inch))
            
            # Результаты анализа
            story.append(Paragraph('Результаты анализа', heading_style))
            
            # Основные показатели
            main_data = [
                ['Показатель', 'Значение'],
                ['Процент пористости:', f'{self.analysis.porosity_percentage:.2f}%' if self.analysis.porosity_percentage else 'Н/Д'],
                ['Количество пор:', str(self.analysis.number_of_pores) if self.analysis.number_of_pores else 'Н/Д'],
                ['Средний размер пор:', f'{self.analysis.average_pore_size:.2f} мкм' if self.analysis.average_pore_size else 'Н/Д'],
                ['Максимальный размер пор:', f'{self.analysis.max_pore_size:.2f} мкм' if self.analysis.max_pore_size else 'Н/Д'],
                ['Минимальный размер пор:', f'{self.analysis.min_pore_size:.2f} мкм' if self.analysis.min_pore_size else 'Н/Д'],
                ['Плотность пор:', f'{self.analysis.pore_density:.4f} пор/мкм²' if self.analysis.pore_density else 'Н/Д'],
                ['Среднее межпоровое расстояние:', f'{self.analysis.average_interpore_distance:.2f} мкм' if self.analysis.average_interpore_distance else 'Н/Д']
            ]
            
            main_table = Table(main_data, colWidths=[3.5*inch, 2.5*inch])
            main_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('FONTNAME', (0, 0), (-1, 0), RUSSIAN_FONT),
                ('FONTSIZE', (0, 0), (-1, 0), 12),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
                ('FONTNAME', (0, 1), (-1, -1), RUSSIAN_FONT),
                ('FONTSIZE', (0, 1), (-1, -1), 11),
                ('GRID', (0, 0), (-1, -1), 1, colors.black)
            ]))
            
            story.append(main_table)
            story.append(PageBreak())
            
            # Добавляем изображения
            story.append(Paragraph('Результаты визуализации', heading_style))
            
            # Получаем список доступных изображений
            available_images = self._get_available_images()
            
            if not available_images:
                story.append(Paragraph('Изображения результатов анализа не найдены.', normal_style))
            else:
                for image_path, description in available_images:
                    try:
                        story.append(Paragraph(description, heading_style))
                        
                        # Проверяем, что файл существует и не пустой
                        if not os.path.exists(image_path) or os.path.getsize(image_path) == 0:
                            story.append(Paragraph(f"[Изображение {description} недоступно]", normal_style))
                            continue
                        
                        # Получаем размеры изображения для правильного масштабирования
                        img = Image.open(image_path)
                        img_width, img_height = img.size
                        aspect = img_height / float(img_width)
                        
                        # Ограничиваем максимальную ширину
                        max_width = 5 * inch
                        if img_width > max_width:
                            display_width = max_width
                            display_height = max_width * aspect
                        else:
                            display_width = img_width
                            display_height = img_height
                        
                        # Добавляем изображение
                        story.append(RLImage(image_path, width=display_width, height=display_height))
                        story.append(Spacer(1, 0.3*inch))
                        
                    except Exception as e:
                        logger.warning(f"Не удалось добавить изображение {image_path}: {e}")
                        story.append(Paragraph(f"[Изображение {description} недоступно]", normal_style))
            
            # Генерируем PDF
            doc.build(story)
            logger.info(f"Отчет PDF сохранен: {output_path}")
            return True
            
        except Exception as e:
            logger.error(f"Ошибка при генерации PDF отчета: {e}")
            return False
    
    def generate_reports(self) -> Dict[str, str]:
        """
        Генерирует отчеты в форматах DOCX и PDF
        
        Returns:
            dict: Словарь с путями к созданным файлам отчетов
        """
        reports = {}
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        
        # Создаем директорию для отчетов если её нет
        reports_dir = os.path.join(self.results_dir, 'reports')
        os.makedirs(reports_dir, exist_ok=True)
        
        # Генерируем DOCX отчет
        docx_path = os.path.join(reports_dir, f'porosity_report_{timestamp}.docx')
        if self.generate_docx_report(docx_path):
            reports['docx'] = docx_path
            logger.info(f"DOCX отчет создан: {docx_path}")
        else:
            logger.error(f"Не удалось создать DOCX отчет: {docx_path}")
        
        # Генерируем PDF отчет
        pdf_path = os.path.join(reports_dir, f'porosity_report_{timestamp}.pdf')
        if self.generate_pdf_report(pdf_path):
            reports['pdf'] = pdf_path
            logger.info(f"PDF отчет создан: {pdf_path}")
        else:
            logger.error(f"Не удалось создать PDF отчет: {pdf_path}")
        
        return reports